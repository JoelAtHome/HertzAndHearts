"""ECG cursor Log Δt snapshots embedded in session Word reports."""

from __future__ import annotations

from datetime import datetime
from pathlib import Path
from tempfile import TemporaryDirectory
import struct
import zlib
import unittest

from docx import Document

from hnh.report import generate_session_report


def _write_tiny_png(path: Path, width: int = 8, height: int = 4) -> None:
    """Write a minimal valid RGB PNG without requiring Pillow."""

    def _chunk(tag: bytes, data: bytes) -> bytes:
        return (
            struct.pack(">I", len(data))
            + tag
            + data
            + struct.pack(">I", zlib.crc32(tag + data) & 0xFFFFFFFF)
        )

    raw = b"".join(b"\x00" + (b"\xff\xff\xff" * width) for _ in range(height))
    ihdr = struct.pack(">IIBBBBB", width, height, 8, 2, 0, 0, 0)
    png = (
        b"\x89PNG\r\n\x1a\n"
        + _chunk(b"IHDR", ihdr)
        + _chunk(b"IDAT", zlib.compress(raw, 9))
        + _chunk(b"IEND", b"")
    )
    path.write_bytes(png)


class ReportEcgCursorCaptureTests(unittest.TestCase):
    def test_report_includes_ecg_cursor_capture_section(self):
        with TemporaryDirectory() as tmp:
            session_dir = Path(tmp)
            img_path = session_dir / "ecg_cursor_capture_20260313_090512.png"
            _write_tiny_png(img_path)

            report_path = session_dir / "session_report.docx"
            data = {
                "session_id": "s-cap-001",
                "profile_id": "Admin",
                "session_type": "General Monitoring",
                "session_start": datetime(2026, 3, 13, 9, 0, 0),
                "session_end": datetime(2026, 3, 13, 9, 15, 0),
                "baseline_hr": 72.0,
                "baseline_rmssd": 28.5,
                "last_hr": 68.0,
                "last_rmssd": 35.1,
                "annotations": [
                    ("09:05:12", "ECG cursor Δt=412.0 ms (QT) (A=1.000s, B=1.412s)"),
                ],
                "ecg_cursor_captures": [
                    {
                        "time": "09:05:12",
                        "dt_ms": 412.0,
                        "interval_type": "QT",
                        "a_t_sec": 1.0,
                        "b_t_sec": 1.412,
                        "image": img_path.name,
                    }
                ],
                "hr_values": [72.0, 70.0, 68.0],
                "rmssd_values": [28.5, 31.0, 35.1],
                "notes": "",
                "csv_path": "session.csv",
                "report_stage": "final",
                "qtc": {},
            }

            generate_session_report(str(report_path), data)
            self.assertTrue(report_path.exists())
            doc = Document(str(report_path))
            text = "\n".join(p.text for p in doc.paragraphs)
            self.assertIn("ECG Cursor Captures", text)
            self.assertIn("09:05:12", text)
            self.assertIn("412.0 ms", text)
            self.assertIn("QT", text)
            self.assertNotIn("(QT)", text)
            image_parts = [
                rel.target_part
                for rel in doc.part.rels.values()
                if "image" in rel.reltype
            ]
            self.assertGreaterEqual(len(image_parts), 1)

    def test_unspecified_interval_omits_type_from_caption(self):
        with TemporaryDirectory() as tmp:
            session_dir = Path(tmp)
            img_path = session_dir / "ecg_cursor_capture_plain.png"
            _write_tiny_png(img_path)
            report_path = session_dir / "session_report.docx"
            data = {
                "session_id": "s-cap-003",
                "profile_id": "Admin",
                "session_type": "General Monitoring",
                "session_start": datetime(2026, 3, 13, 9, 0, 0),
                "session_end": datetime(2026, 3, 13, 9, 15, 0),
                "baseline_hr": 72.0,
                "baseline_rmssd": 28.5,
                "last_hr": 68.0,
                "last_rmssd": 35.1,
                "annotations": [],
                "ecg_cursor_captures": [
                    {
                        "time": "09:06:00",
                        "dt_ms": 1394.3,
                        "interval_type": "",
                        "image": img_path.name,
                    }
                ],
                "hr_values": [72.0],
                "rmssd_values": [28.5],
                "notes": "",
                "csv_path": "session.csv",
                "report_stage": "final",
                "qtc": {},
            }
            generate_session_report(str(report_path), data)
            doc = Document(str(report_path))
            text = "\n".join(p.text for p in doc.paragraphs)
            self.assertIn("09:06:00 — Δt=1394.3 ms", text)
            self.assertNotIn("R-R", text)
            self.assertNotIn("(R-R)", text)

    def test_missing_image_skips_capture_section(self):
        with TemporaryDirectory() as tmp:
            session_dir = Path(tmp)
            report_path = session_dir / "session_report.docx"
            data = {
                "session_id": "s-cap-002",
                "profile_id": "Admin",
                "session_type": "General Monitoring",
                "session_start": datetime(2026, 3, 13, 9, 0, 0),
                "session_end": datetime(2026, 3, 13, 9, 15, 0),
                "baseline_hr": 72.0,
                "baseline_rmssd": 28.5,
                "last_hr": 68.0,
                "last_rmssd": 35.1,
                "annotations": [],
                "ecg_cursor_captures": [
                    {
                        "time": "09:05:12",
                        "dt_ms": 412.0,
                        "interval_type": "QT",
                        "image": "missing.png",
                    }
                ],
                "hr_values": [72.0],
                "rmssd_values": [28.5],
                "notes": "",
                "csv_path": "session.csv",
                "report_stage": "final",
                "qtc": {},
            }
            generate_session_report(str(report_path), data)
            doc = Document(str(report_path))
            text = "\n".join(p.text for p in doc.paragraphs)
            self.assertNotIn("ECG Cursor Captures", text)


if __name__ == "__main__":
    unittest.main()
