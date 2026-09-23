from __future__ import annotations

from datetime import datetime
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest

from docx import Document

from hnh.report import core_metric_reference_text, generate_session_report, generate_session_share_pdf


class ReportReferenceRangeTests(unittest.TestCase):
    def test_core_metric_reference_labels(self):
        self.assertEqual(core_metric_reference_text("hr"), "Personal baseline")
        self.assertEqual(core_metric_reference_text("rmssd"), "Personal baseline")
        self.assertIn("duration-dependent", core_metric_reference_text("sdnn"))
        self.assertEqual(core_metric_reference_text("lf_hf"), "Mid-range zone 1.0–3.0")
        self.assertEqual(core_metric_reference_text("qtc"), "Elevated above 470 ms")
        self.assertIn("80–100 ms", core_metric_reference_text("qrs"))
        self.assertEqual(core_metric_reference_text("unknown"), "—")

    def test_docx_post_session_table_includes_reference_column(self):
        with TemporaryDirectory() as tmp:
            report_path = Path(tmp) / "session_report.docx"
            generate_session_report(str(report_path), _sample_data())
            doc = Document(str(report_path))
            table_text = "\n".join(
                cell.text for table in doc.tables for row in table.rows for cell in row.cells
            )
            self.assertIn("Reference range", table_text)
            self.assertIn("Personal baseline", table_text)
            self.assertIn("Mid-range zone 1.0–3.0", table_text)
            self.assertIn("Elevated above 470 ms", table_text)
            self.assertIn("Typical 80–100 ms; wide above 120 ms", table_text)
            all_text = "\n".join(p.text for p in doc.paragraphs)
            self.assertIn("not a population cutoff", all_text)

    def test_share_pdf_builds_with_reference_column(self):
        with TemporaryDirectory() as tmp:
            pdf_path = Path(tmp) / "session_share.pdf"
            generate_session_share_pdf(str(pdf_path), _sample_data())
            self.assertTrue(pdf_path.exists())
            self.assertGreater(pdf_path.stat().st_size, 1000)


def _sample_data() -> dict:
    return {
        "session_id": "s-001",
        "profile_id": "Admin",
        "session_type": "General Monitoring",
        "session_start": datetime(2026, 2, 25, 9, 0, 0),
        "session_end": datetime(2026, 2, 25, 9, 15, 0),
        "ecg_sensor_name": "Polar H10",
        "baseline_hr": 72.0,
        "baseline_rmssd": 28.5,
        "last_hr": 68.0,
        "last_rmssd": 35.1,
        "annotations": [],
        "hr_values": [72.0, 70.0, 68.0],
        "hrv_values": [40.0, 42.0, 41.0],
        "stress_ratio_values": [1.2, 1.4, 1.1],
        "rmssd_values": [28.5, 31.0, 35.1],
        "notes": "",
        "csv_path": "session.csv",
        "report_stage": "final",
        "qtc": {"session_value_ms": 410, "session_qrs_avg_ms": 92.0},
        "disclaimer": {},
    }


if __name__ == "__main__":
    unittest.main()
